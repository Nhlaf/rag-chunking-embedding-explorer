from pathlib import Path
import re

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".xls"}


def _parse_pdf(path: Path) -> tuple[str, int, list[dict]]:
    reader = PdfReader(str(path))
    text = "\n\n".join(page.extract_text() or "" for page in reader.pages).strip()
    return text, len(reader.pages), []


def _parse_docx(path: Path) -> tuple[str, int | None, list[dict]]:
    document = Document(str(path))
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    tables = []
    lines = list(paragraphs)
    for table_index, table in enumerate(document.tables, 1):
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        rows = [row for row in rows if any(row)]
        if rows:
            tables.append({"name": f"Table {table_index}", "rows": rows})
            lines.append(f"Table {table_index}:")
            lines.extend(" | ".join(value for value in row if value) for row in rows)
    return "\n\n".join(lines).strip(), None, tables


def _parse_excel(path: Path) -> tuple[str, int, list[dict]]:
    if path.suffix.lower() == ".xls":
        import xlrd
        workbook = xlrd.open_workbook(path)
        sheets = []
        for sheet in workbook.sheets():
            rows = [sheet.row_values(index) for index in range(sheet.nrows)]
            sheets.append({"name": sheet.name, "rows": rows})
    else:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheets = [{"name": sheet.title, "rows": [list(row) for row in sheet.iter_rows(values_only=True)]} for sheet in workbook.worksheets]
    lines = []
    for sheet in sheets:
        lines.append(f"Sheet: {sheet['name']}")
        rows = sheet["rows"]
        headers = [str(value) if value is not None else f"Column {index + 1}" for index, value in enumerate(rows[0])] if rows else []
        for row in rows[1:] if rows else []:
            values = [f"{headers[index]}: {value}" for index, value in enumerate(row) if value is not None]
            if values:
                lines.append(". ".join(values) + ".")
    return "\n".join(lines).strip(), len(sheets), sheets


def parse_document(path: Path) -> dict:
    extension = path.suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError("Unsupported file type. Use PDF, DOCX, XLSX, or XLS.")
    if extension == ".pdf":
        text, count, sheets = _parse_pdf(path)
    elif extension == ".docx":
        text, count, sheets = _parse_docx(path)
    else:
        text, count, sheets = _parse_excel(path)
    if not text.strip():
        raise ValueError("The document did not contain extractable text.")
    extracted_json = []
    for sheet in sheets:
        rows = sheet.get("rows", [])
        if not rows:
            continue
        headers = [str(value).strip() if value is not None else f"Column {index + 1}" for index, value in enumerate(rows[0])]
        for row in rows[1:]:
            extracted_json.append({headers[index]: value for index, value in enumerate(row) if index < len(headers) and value is not None})
    structured = _looks_structured(text, sheets)
    if structured and not extracted_json:
        extracted_json = [{"heading": line} for line in text.splitlines() if re.match(r"^\d+(?:\.\d+)*\s+[^.]{2,80}$", line.strip())]
    if sheets and any(sheet.get("rows") for sheet in sheets):
        reason = "The content has repeated table rows or records."
    elif structured:
        reason = "The content is organized into repeated numbered sections, so it is structured document data."
    else:
        reason = "The content is mainly narrative paragraphs without a consistent row-and-column or repeated-section pattern."
    if not structured:
        reason = "The content is mainly narrative paragraphs without a consistent row-and-column or repeated-section pattern."
    return {
        "extracted_text": text,
        "page_count": count,
        "sheets": sheets,
        "document_type": "structured" if structured else "unstructured",
        "extracted_json": extracted_json,
        "classification_reason": reason,
    }


def _looks_structured(text: str, tables: list[dict]) -> bool:
    if tables:
        return True
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if len(lines) < 2:
        return False
    delimited = [line for line in lines if line.count("|") >= 1 or line.count("\t") >= 1]
    if len(delimited) >= 2:
        widths = [len(re.split(r"\||\t", line)) for line in delimited]
        return len(set(widths)) == 1
    numbered_headings = [line for line in lines if re.match(r"^\d+(?:\.\d+)*\s+[^.]{2,80}$", line)]
    if len(numbered_headings) >= 2:
        return True
    key_value_lines = [line for line in lines if re.match(r"^[^:]{1,40}:\s+.+", line)]
    return len(key_value_lines) >= 3
